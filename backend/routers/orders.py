import io
from datetime import datetime, timezone
from typing import List
from zoneinfo import ZoneInfo

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pymongo.errors import DuplicateKeyError

from lib.db import db
from lib.stock import consume_for_items, restore_moves
from models.order import NextNumber, Order, OrderInput, OrderSaved, PaidUpdate

router = APIRouter()

TYPE_LABELS = {"revenda": "Revenda", "cliente_final": "Cliente Final"}
TZ = ZoneInfo("America/Sao_Paulo")


def _local(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).astimezone(TZ).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return iso


@router.get("/orders", response_model=List[Order])
async def list_orders():
    docs = await db.orders.find({}, {"_id": 0}).sort("number", -1).to_list(5000)
    return [Order(**d) for d in docs]


@router.get("/orders/next-number", response_model=NextNumber)
async def next_number():
    # Next = last saved order's number + 1 (skipping numbers already taken).
    last = await db.orders.find_one({}, {"_id": 0, "number": 1}, sort=[("created_at", -1)])
    n = (last["number"] + 1) if last else 1
    while await db.orders.find_one({"number": n}, {"_id": 1}):
        n += 1
    return NextNumber(next_number=n)


@router.get("/orders/export")
async def export_orders():
    docs = await db.orders.find({}, {"_id": 0}).sort("number", -1).to_list(10000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pedidos"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="3D2314")
    ws.append(["Pedido Nº", "Data", "Cliente", "Tipo", "Itens", "Total (R$)",
               "Forma de Pagamento", "Prazo p/ Pagamento", "Entrega", "Observações", "Quitado"])
    ws2 = wb.create_sheet("Itens")
    ws2.append(["Pedido Nº", "Data", "Cliente", "Tipo", "SKU", "Produto",
                "Quantidade", "Preço Unit. (R$)", "Subtotal (R$)"])
    for sheet in (ws, ws2):
        for c in sheet[1]:
            c.font, c.fill = head_font, head_fill
    for d in docs:
        o = Order(**d)
        date = _local(o.created_at)
        tipo = TYPE_LABELS.get(o.customer_type, o.customer_type)
        items_txt = "; ".join(
            f"{i.quantity}x {i.name}" + (f" [{i.sku}]" if i.sku else "") for i in o.items
        )
        ws.append([o.number, date, o.customer_name, tipo, items_txt, round(o.total, 2),
                   o.payment_method, o.payment_term, o.delivery, o.notes, "Sim" if o.paid else "Não"])
        for i in o.items:
            ws2.append([o.number, date, o.customer_name, tipo, i.sku, i.name, i.quantity,
                        round(i.price, 2), round(i.price * i.quantity, 2)])
    for sheet, widths in ((ws, [11, 17, 28, 14, 60, 12, 20, 20, 22, 30, 10]),
                          (ws2, [11, 17, 28, 14, 14, 36, 11, 15, 14])):
        for idx, w in enumerate(widths):
            sheet.column_dimensions[chr(65 + idx)].width = w

    # Packaging: current stock + full movement log.
    kind_labels = {"inicial": "Estoque inicial", "entrada": "Entrada", "ajuste": "Ajuste manual",
                   "pedido": "Baixa por pedido", "devolucao": "Devolução (pedido excluído)"}
    packs = await db.packaging.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    ws3 = wb.create_sheet("Estoque Embalagens")
    ws3.append(["Embalagem", "Quantidade", "Estoque Mínimo", "Situação"])
    for p in packs:
        q, mn = p["quantity"], p.get("min_quantity", 0)
        ws3.append([p["name"], q, mn, "Sem estoque" if q <= 0 else "Baixo" if q <= mn else "OK"])
    logs = await db.stock_log.find({}, {"_id": 0}).sort("created_at", -1).to_list(50000)
    ws4 = wb.create_sheet("Movimentações Estoque")
    ws4.append(["Data", "Embalagem", "Tipo", "Quantidade", "Saldo", "Pedido Nº"])
    for lg in logs:
        ws4.append([_local(lg["created_at"]), lg["packaging_name"], kind_labels.get(lg["kind"], lg["kind"]),
                    lg["delta"], lg["balance"], lg.get("order_number")])
    for sheet, widths in ((ws3, [24, 12, 15, 13]), (ws4, [17, 24, 28, 12, 10, 11])):
        for c in sheet[1]:
            c.font, c.fill = head_font, head_fill
        for idx, w in enumerate(widths):
            sheet.column_dimensions[chr(65 + idx)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"pedidos-artezano-{datetime.now(TZ).strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/orders/{id}", response_model=Order)
async def get_order(id: str):
    doc = await db.orders.find_one({"id": id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Pedido não encontrado")
    return Order(**doc)


@router.post("/orders", response_model=OrderSaved)
async def create_order(data: OrderInput):
    total = round(sum(i.price * i.quantity for i in data.items), 2)
    order = Order(**data.model_dump(), total=total)
    try:
        await db.orders.insert_one(order.model_dump())
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail=f"O pedido nº {data.number} já existe")
    moves, warnings = await consume_for_items(data.items, data.number)
    if moves:
        await db.orders.update_one({"id": order.id}, {"$set": {"stock_moves": moves}})
    return OrderSaved(**order.model_dump(exclude={"stock_moves"}), stock_moves=moves, stock_warnings=warnings)


@router.patch("/orders/{id}/paid", response_model=Order)
async def set_paid(id: str, data: PaidUpdate):
    paid_at = datetime.now(timezone.utc).isoformat() if data.paid else None
    doc = await db.orders.find_one_and_update(
        {"id": id}, {"$set": {"paid": data.paid, "paid_at": paid_at}},
        projection={"_id": 0}, return_document=True,
    )
    if not doc:
        raise HTTPException(status_code=404, detail="Pedido não encontrado")
    return Order(**doc)


@router.delete("/orders/{id}")
async def delete_order(id: str):
    doc = await db.orders.find_one_and_delete({"id": id}, projection={"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Pedido não encontrado")
    await restore_moves(doc.get("stock_moves", []), doc["number"])
    return {"ok": True}
